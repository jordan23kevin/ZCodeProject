# -*- coding: utf-8 -*-
import json
import os
import sqlite3
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

CST = timezone(timedelta(hours=8))
PRODUCT_DB = r"E:\Kimi Code\temu分析\报活动\base_data.db"
REPORT_ROOT = r"E:\Temu账单\reports"


def money(value):
    return round(float(value or 0), 2)


def cost_by_month(con, product_db=PRODUCT_DB):
    latest = con.execute("SELECT MAX(file_key) FROM settlement_eu").fetchone()[0]
    if not latest:
        raise RuntimeError("没有欧区结算数据，无法计算成本")
    con.execute("ATTACH DATABASE ? AS products", (f"file:{product_db}?mode=ro",))
    con.executescript("""
    CREATE TEMP TABLE IF NOT EXISTS recon_bill_rows AS
    SELECT file_key,po,sku_id,sku_code,qty,'settled' src FROM settlement_eu WHERE tx_type='销售回款'
    UNION ALL
    SELECT file_key,po,sku_id,sku_code,qty,'pending' FROM payment WHERE status='待处理';
    CREATE TEMP TABLE IF NOT EXISTS recon_sku_ids AS SELECT DISTINCT CAST(sku_id AS INTEGER) sku_id FROM recon_bill_rows;
    CREATE TEMP TABLE IF NOT EXISTS recon_sku_map AS
    SELECT i.sku_id,GROUP_CONCAT(DISTINCT b.skc_code) codes
    FROM recon_sku_ids i LEFT JOIN products.base_data b ON b.sku_id=i.sku_id GROUP BY i.sku_id;
    """)
    rows = con.execute("""
        SELECT file_key,src,qty,sku_code,codes FROM recon_bill_rows r
        LEFT JOIN recon_sku_map m ON m.sku_id=CAST(r.sku_id AS INTEGER)
    """).fetchall()
    result = {}
    for file_key, src, qty, sku_code, mapped in rows:
        month = file_key.split("|")[-1]
        code = str(sku_code or mapped or "").upper()
        unit_cost = 58 if code.startswith(("HX", "RX")) else 33
        result.setdefault(month, {"settled": 0.0, "pending": 0.0})
        result[month][src] += (qty or 0) * unit_cost
    return result


def month_metrics(con, month, costs):
    actual = con.execute(
        "SELECT COALESCE(SUM(amount),0) FROM ledger WHERE file_key=? AND tx_type NOT IN('充值','提现')",
        (f"ledger|{month}",),
    ).fetchone()[0]
    pending = con.execute(
        "SELECT COALESCE(SUM(sale_back+ship_back+sale_reversal+ship_reversal),0) FROM payment WHERE file_key=? AND status='待处理'",
        (f"payment_pending|{month}",),
    ).fetchone()[0]
    return_fee = con.execute(
        "SELECT COALESCE(SUM(r.amount),0) FROM return_label r JOIN payment p ON p.po=r.po AND p.status='待处理' WHERE p.file_key=?",
        (f"payment_pending|{month}",),
    ).fetchone()[0]
    ship_fee = con.execute(
        "SELECT COALESCE(SUM(s.amount),0) FROM ship_label s JOIN orders o ON o.tracking_no=s.waybill_no JOIN payment p ON p.po=o.order_no AND p.status='待处理' WHERE p.file_key=?",
        (f"payment_pending|{month}",),
    ).fetchone()[0]
    estimated = 0.0
    if con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='return_fee_estimate'").fetchone():
        estimated = con.execute(
            "SELECT COALESCE(-SUM(amount_cny),0) FROM return_fee_estimate WHERE status='open' AND substr(created_at,1,7)=?",
            (month,),
        ).fetchone()[0]
    settled_cost = costs.get(month, {}).get("settled", 0.0)
    pending_cost = costs.get(month, {}).get("pending", 0.0)
    pending_direct = return_fee + ship_fee
    settled_profit = actual - settled_cost - pending_direct + estimated
    pending_profit = pending - pending_cost + pending_direct
    return {
        "month": month,
        "actual_net": money(actual),
        "pending_net": money(pending),
        "settled_cost": money(settled_cost),
        "pending_cost": money(pending_cost),
        "pending_direct_fees": money(pending_direct),
        "estimated_fees": money(estimated),
        "settled_profit": money(settled_profit),
        "pending_profit": money(pending_profit),
        "total_profit": money(settled_profit + pending_profit),
    }


def ensure_history_schema(con):
    con.executescript("""
    CREATE TABLE IF NOT EXISTS recon_run (
      run_id TEXT PRIMARY KEY, created_at TEXT, period_start TEXT, period_end TEXT,
      settled_profit REAL, pending_profit REAL, total_profit REAL,
      report_path TEXT, status TEXT, error TEXT);
    CREATE TABLE IF NOT EXISTS recon_month (
      run_id TEXT NOT NULL, month TEXT NOT NULL,
      actual_net REAL, pending_net REAL, settled_cost REAL, pending_cost REAL,
      pending_direct_fees REAL, estimated_fees REAL,
      settled_profit REAL, pending_profit REAL, total_profit REAL,
      PRIMARY KEY(run_id,month));
    """)


def collect_report(db_path, start=None, end=None, product_db=PRODUCT_DB):
    con = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    try:
        costs = cost_by_month(con, product_db)
        months = sorted(set(costs) | {row[0].split("|")[-1] for row in con.execute("SELECT DISTINCT file_key FROM ledger")})
        if start:
            months = [month for month in months if month >= start[:7]]
        if end:
            months = [month for month in months if month <= end[:7]]
        reports = [month_metrics(con, month, costs) for month in months]
        totals = {key: money(sum(row[key] for row in reports)) for key in reports[0] if key != "month"} if reports else {}
        totals["month"] = "合计"
        notes = []
        pending_overlap = con.execute(
            "SELECT COUNT(*) FROM payment p JOIN payment r ON p.po=r.po AND p.sku_id=r.sku_id WHERE p.status='待处理' AND r.status='已到账'"
        ).fetchone()[0]
        if pending_overlap:
            notes.append({"level": "error", "title": "已到账与待处理重叠", "detail": f"{pending_overlap} 条 PO+SKU 需要按金额复核"})
        else:
            notes.append({"level": "ok", "title": "款项状态去重", "detail": "已到账与待处理没有重复 PO+SKU"})
        for row in con.execute("SELECT file_key FROM ledger GROUP BY file_key"):
            month = row[0].split("|")[-1]
            ledger_settle = con.execute("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE file_key=? AND tx_type='结算'", (row[0],)).fetchone()[0]
            eu_settle = con.execute("SELECT COALESCE(SUM(amount),0) FROM settlement_eu WHERE file_key=?", (f"europe|{month}",)).fetchone()[0]
            difference = money(ledger_settle - eu_settle)
            if abs(difference) > 0.01:
                notes.append({"level": "warn", "title": f"{month} 结算差异", "detail": f"卖家中心与欧区相差 {difference:+.2f} 元"})
        pending_total = con.execute("SELECT COALESCE(SUM(sale_back+ship_back+sale_reversal+ship_reversal),0) FROM payment WHERE status='待处理'").fetchone()[0]
        notes.append({"level": "info", "title": "未出账款项", "detail": f"当前待处理净额 {pending_total:.2f} 元；未发生退款不做预估"})
        estimate_total = 0.0
        if con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='return_fee_estimate'").fetchone():
            estimate_total = con.execute("SELECT COALESCE(SUM(amount_cny),0) FROM return_fee_estimate WHERE status='open'").fetchone()[0]
        notes.append({"level": "info", "title": "退货费暂估", "detail": f"有效站点暂估 {estimate_total:.2f} 元；实际面单出现后完整回正再按实际入账"})
        _vio_cols = {row[1] for row in con.execute("PRAGMA table_info(violation)")}
        _appeal_col = "appeal_status" if "appeal_status" in _vio_cols else "NULL"
        latest_violations = con.execute(f"SELECT violation_id,po,vtype,amount,tx_time,{_appeal_col} FROM violation ORDER BY tx_time DESC LIMIT 5").fetchall()
        for row in latest_violations:
            notes.append({"level": "warn", "title": f"{row[2]} {row[3]:.2f} 元", "detail": f"{row[0]} | {row[1]} | {row[4]} | 处理状态 {row[5] or '未记录'}"})
        sources = {
            "ledger_latest": con.execute("SELECT MAX(tx_time) FROM ledger").fetchone()[0],
            "settlement_latest": con.execute("SELECT MAX(arrive_time) FROM settlement_eu").fetchone()[0],
            "orders_latest": con.execute("SELECT MAX(created_at) FROM orders").fetchone()[0],
            "aftersale_latest": con.execute("SELECT MAX(create_time) FROM aftersale").fetchone()[0],
            "pending_rows": con.execute("SELECT COUNT(*) FROM payment WHERE status='待处理'").fetchone()[0],
            "settled_rows": con.execute("SELECT COUNT(*) FROM payment WHERE status='已到账'").fetchone()[0],
            "order_rows": con.execute("SELECT COUNT(*) FROM orders").fetchone()[0],
            "aftersale_rows": con.execute("SELECT COUNT(*) FROM aftersale").fetchone()[0],
        }
        return {"months": reports, "totals": totals, "notes": notes, "sources": sources}
    finally:
        con.close()


def write_excel(report, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "月度汇总"
    headers = ["月份", "已到账净额(不含划拨)", "未到账净额", "已到账胚衣成本", "未到账胚衣成本", "未到账直接费用", "退货费暂估", "已到账利润", "未到账利润", "总利润"]
    ws.append(headers)
    keys = ["month", "actual_net", "pending_net", "settled_cost", "pending_cost", "pending_direct_fees", "estimated_fees", "settled_profit", "pending_profit", "total_profit"]
    for row in report["months"] + [report["totals"]]:
        ws.append([row[key] for key in keys])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    widths = [12, 24, 16, 18, 18, 18, 14, 16, 16, 14]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = '#,##0.00;[Red]-#,##0.00'
    ws2 = wb.create_sheet("财务注意事项")
    ws2.append(["级别", "事项", "说明"])
    for note in report["notes"]:
        ws2.append([note["level"], note["title"], note["detail"]])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    ws2.column_dimensions["A"].width = 10
    ws2.column_dimensions["B"].width = 32
    ws2.column_dimensions["C"].width = 110
    ws3 = wb.create_sheet("数据来源")
    ws3.append(["项目", "值"])
    for key, value in report["sources"].items():
        ws3.append([key, value])
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 50
    wb.save(path)


def generate_report(db_path, start=None, end=None, report_root=REPORT_ROOT, product_db=PRODUCT_DB):
    report = collect_report(db_path, start, end, product_db)
    run_id = f"{datetime.now(CST):%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}"
    directory = Path(report_root) / run_id
    directory.mkdir(parents=True, exist_ok=True)
    xlsx_path = directory / f"Temu对账_{run_id}.xlsx"
    json_path = directory / f"Temu对账_{run_id}.json"
    write_excel(report, xlsx_path)
    payload = {"run_id": run_id, "created_at": datetime.now(CST).isoformat(timespec="seconds"), **report, "report_path": str(xlsx_path)}
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    con = sqlite3.connect(db_path)
    try:
        ensure_history_schema(con)
        with con:
            con.execute(
                "INSERT INTO recon_run(run_id,created_at,period_start,period_end,settled_profit,pending_profit,total_profit,report_path,status) VALUES(?,?,?,?,?,?,?,?,?)",
                (run_id, payload["created_at"], start, end, report["totals"].get("settled_profit"), report["totals"].get("pending_profit"), report["totals"].get("total_profit"), str(xlsx_path), "completed"),
            )
            for row in report["months"]:
                con.execute(
                    "INSERT INTO recon_month VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                    (run_id, row["month"], row["actual_net"], row["pending_net"], row["settled_cost"], row["pending_cost"], row["pending_direct_fees"], row["estimated_fees"], row["settled_profit"], row["pending_profit"], row["total_profit"]),
                )
    finally:
        con.close()
    return payload


def history(db_path, limit=20):
    con = sqlite3.connect(db_path)
    try:
        ensure_history_schema(con)
        rows = con.execute(
            "SELECT run_id,created_at,period_start,period_end,settled_profit,pending_profit,total_profit,report_path,status,error FROM recon_run ORDER BY created_at DESC LIMIT ?",
            (limit,),
        ).fetchall()
        keys = ["run_id", "created_at", "period_start", "period_end", "settled_profit", "pending_profit", "total_profit", "report_path", "status", "error"]
        return [dict(zip(keys, row)) for row in rows]
    finally:
        con.close()
